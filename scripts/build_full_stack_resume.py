from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


WORKSPACE_ROOT = Path(__file__).resolve().parents[2]
OUTPUT = WORKSPACE_ROOT / "output/docx/Neeraj_Sharma_Full_Stack_Engineering_Resume.docx"


def set_font(run, size=9, bold=False, color="202124"):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def remove_paragraph_borders(element):
    properties = element.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is not None:
        properties.remove(borders)


def add_hyperlink(paragraph, text, url):
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    node = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    properties.append(fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    properties.append(size)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2446A8")
    properties.append(color)
    node.append(properties)
    value = OxmlElement("w:t")
    value.text = text
    node.append(value)
    hyperlink.append(node)
    paragraph._p.append(hyperlink)


def add_section(doc, label):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True
    set_font(paragraph.add_run(label.upper()), size=10, bold=True, color="000000")
    return paragraph


def add_role(doc, title, organization, dates, bullets):
    heading = doc.add_paragraph()
    heading.paragraph_format.space_before = Pt(3)
    heading.paragraph_format.space_after = Pt(1)
    heading.paragraph_format.keep_with_next = True
    heading.paragraph_format.tab_stops.add_tab_stop(Inches(7.38), WD_TAB_ALIGNMENT.RIGHT)
    set_font(heading.add_run(f"{title} | {organization}"), size=9.5, bold=True, color="000000")
    set_font(heading.add_run(f"\t{dates}"), size=9, bold=True, color="333333")

    for text in bullets:
        paragraph = doc.add_paragraph(style="Resume Bullet")
        paragraph.paragraph_format.keep_together = True
        set_font(paragraph.add_run("• " + text), size=9.1)


def add_project(doc, name, description):
    paragraph = doc.add_paragraph(style="Resume Bullet")
    paragraph.paragraph_format.keep_together = True
    set_font(paragraph.add_run(f"• {name}: "), size=9.1, bold=True, color="000000")
    set_font(paragraph.add_run(description), size=9.1)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.4)
    section.bottom_margin = Inches(0.4)
    section.left_margin = Inches(0.56)
    section.right_margin = Inches(0.56)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(8.8)
    normal.font.color.rgb = RGBColor(32, 33, 36)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    bullet = doc.styles.add_style("Resume Bullet", WD_STYLE_TYPE.PARAGRAPH)
    bullet.base_style = normal
    bullet.paragraph_format.left_indent = Inches(0.16)
    bullet.paragraph_format.first_line_indent = Inches(-0.12)
    bullet.paragraph_format.space_after = Pt(0.7)
    bullet.paragraph_format.line_spacing = 1.0

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(0)
    title.paragraph_format.keep_with_next = True
    remove_paragraph_borders(title.style.element)
    remove_paragraph_borders(title._p)
    set_font(title.add_run("Neeraj Kumar Sharma"), size=21, bold=True, color="000000")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(2)
    subtitle.paragraph_format.keep_with_next = True
    set_font(
        subtitle.add_run("FULL-STACK ENGINEER | HANDS-ON TECH LEAD | AI AND SAAS"),
        size=9.8,
        bold=True,
        color="000000",
    )

    contact = doc.add_paragraph()
    contact.paragraph_format.space_after = Pt(4)
    contact.paragraph_format.keep_with_next = True
    add_hyperlink(contact, "ittsneeraj@gmail.com", "mailto:ittsneeraj@gmail.com")
    set_font(contact.add_run("  |  +91 7470915225  |  Bengaluru, India  |  "), size=8.5, color="444444")
    add_hyperlink(contact, "itsneeraj.com", "https://itsneeraj.com")
    set_font(contact.add_run("  |  "), size=8.5, color="444444")
    add_hyperlink(contact, "LinkedIn", "https://linkedin.com/in/neer4j")
    set_font(contact.add_run("  |  "), size=8.5, color="444444")
    add_hyperlink(contact, "GitHub", "https://github.com/NEER4J")

    summary = doc.add_paragraph()
    summary.paragraph_format.space_after = Pt(2)
    summary.paragraph_format.line_spacing = 1.03
    set_font(
        summary.add_run(
            "Full-stack engineer with 5+ years building AI and SaaS products across frontend, backend, data, and integrations. "
            "Hands-on in implementation, with technical leadership across architecture and delivery. Shipped platforms serving "
            "2,000+ users and delivering 100k+ messages; reduced client support workload by 40%."
        ),
        size=9.3,
    )

    impact = doc.add_paragraph()
    impact.paragraph_format.space_after = Pt(1)
    set_font(impact.add_run("IMPACT  "), size=8.8, bold=True, color="000000")
    set_font(
        impact.add_run("2,000+ users  |  300+ grants refreshed daily  |  100k+ messages  |  40% less support workload  |  $5k+/mo client savings"),
        size=8.8,
        color="333333",
    )

    add_section(doc, "Technical Skills")
    skills = [
        ("Languages", "TypeScript, JavaScript, Python, SQL"),
        ("Frontend", "Next.js, React, Tailwind CSS, Shadcn UI"),
        ("Backend and data", "Node.js, Express, Django, REST APIs, webhooks, PostgreSQL, Supabase, MongoDB"),
        ("AI and infrastructure", "OpenAI, Claude, Gemini, RAG, LangChain, Vercel AI SDK, n8n, Docker, Vercel, Railway"),
        ("Systems and leadership", "Multi-tenant SaaS, AI workflows, payments, analytics, architecture, technical direction, delivery ownership"),
    ]
    for label, values in skills:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0.3)
        paragraph.paragraph_format.keep_together = True
        set_font(paragraph.add_run(f"{label}: "), size=8.8, bold=True, color="000000")
        set_font(paragraph.add_run(values), size=8.8)

    add_section(doc, "Engineering Experience")
    add_role(
        doc,
        "Founder and Lead Engineer",
        "Docsiv",
        "06/2026-Present",
        [
            "Designed and built a multi-tenant AI document workspace in Next.js, TypeScript, and Supabase, spanning editors, brand kits, client portals, collaboration, analytics, signing, and credit billing.",
            "Own technical direction from architecture through release; launched the product and continue shipping improvements with early agency users.",
        ],
    )
    add_role(
        doc,
        "Lead Engineer, AI Products",
        "Virtual Xcellence",
        "01/2026-Present",
        [
            "Govgrant.ca: Built RAG-based grant matching for 2,000+ users and an automated pipeline refreshing 300+ grants daily, alongside authentication, subscriptions, billing, and administration.",
            "SpeedIQ: Led engineering delivery for a multi-tenant WhatsApp and email platform with Meta Business API onboarding, broadcasts, chatbots, live chat, and analytics; delivered 100k+ messages.",
            "Set technical direction across AI workflows, integrations, and platform systems while balancing delivery, data freshness, and production reliability.",
        ],
    )
    add_role(
        doc,
        "Full-Stack Developer",
        "NJ Designpark",
        "10/2022-01/2026",
        [
            "Built and shipped multi-tenant SaaS for education, trades, and service businesses, owning implementation across application, APIs, data, and deployment.",
            "Integrated OpenAI and Claude into production workflows, reducing client support workload by 40%.",
            "Delivered custom business systems, monitoring, and automation credited with saving clients more than $5,000 per month; coordinated technical delivery with client teams.",
        ],
    )
    add_role(
        doc,
        "Independent Full-Stack Developer",
        "Freelance",
        "2020-2022",
        [
            "Delivered production web applications and payment workflows for small businesses using Django, Next.js, Node.js, Stripe, Razorpay, and PayPal.",
        ],
    )

    add_section(doc, "Selected Systems")
    add_project(
        doc,
        "Habiv",
        "Built creator and player workflows, game ingestion and conversion, storage, analytics, leaderboards, and isolated game execution for a browser-first game platform.",
    )
    add_project(
        doc,
        "Apstic",
        "Shipped AI automation across CRM, commerce, accounting, messaging, and browser workflows, plus a local-first assistant.",
    )

    add_section(doc, "Education")
    education = doc.add_paragraph()
    education.paragraph_format.keep_together = True
    set_font(
        education.add_run("B.Tech in Computer Science | Chhattisgarh Swami Vivekanand Technical University | 2020-2023"),
        size=9,
        bold=True,
        color="000000",
    )

    properties = doc.core_properties
    properties.title = "Neeraj Sharma Full-Stack Engineering Resume"
    properties.subject = "Full-stack engineering and hands-on technical leadership"
    properties.author = "Neeraj Kumar Sharma"
    properties.keywords = "Full-Stack Engineer, Technical Lead, AI, SaaS, RAG, Next.js, TypeScript"
    properties.comments = ""

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
